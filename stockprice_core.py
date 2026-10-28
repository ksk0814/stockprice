# -*- coding: utf-8 -*-
"""
stockprice_core.py
理論株価の計算（残余利益モデル）を Web アプリから呼べるよう関数化したもの。

- CSV (IR BANK 形式の fy-balance-sheet.csv / fy-profit-and-loss.csv) の読み込み
- 銘柄コードから IR BANK の CSV を取得（失敗時は Yahoo Finance の財務データで代替）
- Yahoo Finance から株価・流動資産を取得
- 残余利益モデル（Ohlson 1995 ほか）による理論株価、感応度、株価との乖離、参考指標（PBR・PER・PSR・正味流動資産）
- 計算式入りの Excel バイト列の生成
"""

import csv
import datetime as dt
import io
import os
import re
from dataclasses import dataclass, field

import openpyxl as px
import requests

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"
}

UPSIDE_THRESHOLD = 30   # 株価と理論株価の乖離（%）がこれを超えたら割安圏・割高圏（Graham の安全余裕率の考え方）
COST_OF_EQUITY = 0.08   # 株主資本コストの既定値（伊藤レポート 2014 が示した ROE の最低ライン 8%）
FADE_YEARS = 10         # ROE が株主資本コストに収れんするまでの年数
DEFAULT_PAYOUT = 0.30   # 配当性向を取得できないときの仮定

# 証券コード（例: 7203, 130A）。"2024/03" のような年度表記は除外する
CODE_RE = re.compile(r"(?<![0-9A-Za-z/])([0-9][0-9A-Z]{2}[0-9A-Z])(?![0-9A-Za-z/])")


# ============================================================
# 共通ユーティリティ
# ============================================================
def safe_float(val):
    """文字列・数値・None が混在するデータから安全に float を取得する（取得できなければ None）"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(",", "").replace("¥", "").replace("$", "").replace("%", "")
    if s in ("", "-", "－", "—"):
        return None
    s = re.sub(r"[^0-9.\-]", "", s)
    try:
        return float(s) if s not in ("", "-", ".") else None
    except ValueError:
        return None


def decode_csv(data):
    """CSV のバイト列を文字列に変換（UTF-8 / BOM付き / Shift_JIS に対応）"""
    if isinstance(data, str):
        return data
    for enc in ("utf-8-sig", "cp932"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def read_csv_rows(data):
    """CSV を行のリスト（各行は文字列リスト）として返す。空行も保持する（元スクリプトの行番号と合わせるため）"""
    text = decode_csv(data)
    return [row for row in csv.reader(io.StringIO(text), delimiter=",")]


def _last_row_index(rows):
    """A 列が空でない最終行のインデックス（元スクリプトの maxRow 判定と同じ）"""
    for i in reversed(range(len(rows))):
        if rows[i] and rows[i][0].strip() != "":
            return i
    raise ValueError("CSV にデータ行がありません。")


def _header_index(rows, names, default):
    """見出し行から列番号を探す。見つからなければ元スクリプトの固定列番号を使う"""
    for row in rows[:10]:
        for idx, cell in enumerate(row):
            if cell.strip() in names:
                return idx
    return default


def _cell(rows, r, c):
    try:
        return rows[r][c]
    except IndexError:
        return None


# ============================================================
# 入力データ
# ============================================================
@dataclass
class FinancialInputs:
    code: str = ""
    company: str = ""           # Excel の A1 / ファイル名に使う表示名
    bps: float = None           # D2
    equity_ratio: float = None  # F2（小数。例 0.45）
    eps: float = None           # D4
    roa: float = None           # E4（小数）
    roe_hist: list = field(default_factory=list)  # 過去の実績 ROE（小数、古い順）
    payout_ratio: float = None  # 配当性向（小数）。取得できなければ None
    sales: float = None         # G12（円）
    total_assets: float = None  # 円
    net_assets: float = None    # 円
    shareholders_equity: float = None  # 円
    current_assets_million: float = None  # G14（百万円）
    stock_price: float = None   # C8（円）
    shares_outstanding: float = None
    source: str = ""
    raw_pl: bytes = None        # 取得した PL CSV（IR レポートで再利用）
    notes: list = field(default_factory=list)


def parse_irbank_csv(bs_data, pl_data):
    """IR BANK 形式の BS/PL CSV から指標を取り出す（StockPrice_ver2.py のセル3前半と同じ列）"""
    bs = read_csv_rows(bs_data)
    pl = read_csv_rows(pl_data)
    inp = FinancialInputs(source="CSV")

    # ---- BS ----
    r = _last_row_index(bs)
    c_total = _header_index(bs, ("総資産",), 1)
    c_net = _header_index(bs, ("純資産",), 2)
    c_equity = _header_index(bs, ("株主資本",), 3)
    c_bps = _header_index(bs, ("BPS",), 7)
    c_ratio = _header_index(bs, ("自己資本比率",), 8)

    inp.bps = safe_float(_cell(bs, r, c_bps))
    ratio = safe_float(_cell(bs, r, c_ratio))
    inp.equity_ratio = ratio / 100 if ratio is not None else None
    inp.total_assets = safe_float(_cell(bs, r, c_total))
    inp.net_assets = safe_float(_cell(bs, r, c_net))
    inp.shareholders_equity = safe_float(_cell(bs, r, c_equity))
    if inp.bps is None:
        raise ValueError(f"BS の最終行（{r + 1} 行目）から BPS を読み取れません。ファイルを確認してください。")
    if inp.equity_ratio is None and inp.shareholders_equity and inp.total_assets:
        inp.equity_ratio = inp.shareholders_equity / inp.total_assets

    # ---- PL ----
    m = _last_row_index(pl)
    c_sales = _header_index(pl, ("売上高", "営業収益", "経常収益"), 1)
    c_eps = _header_index(pl, ("EPS",), 5)
    c_roa = _header_index(pl, ("ROA",), 7)
    c_roe = _header_index(pl, ("ROE",), 6)

    is_forecast = any("予想" in str(v) for v in pl[m])
    if is_forecast:
        inp.notes.append(f"PL の最終行（{m + 1} 行目）は予想データです。EPS は予想値を使います。")

    inp.eps = safe_float(_cell(pl, m, c_eps))
    roa = safe_float(_cell(pl, m, c_roa))
    if roa is None and m - 1 >= 0:
        roa = safe_float(_cell(pl, m - 1, c_roa))
        if roa is not None:
            inp.notes.append("最終行に ROA が無いため、1 つ前の行の ROA を使いました。")
    inp.roa = roa / 100 if roa is not None else None

    # 売上高：元スクリプトと同じく maxRow-1 行目（予想行がある前提の最新実績行）
    sales_row = m - 1 if is_forecast else m
    inp.sales = safe_float(_cell(pl, sales_row, c_sales))
    if inp.sales is None and sales_row != m - 1:
        inp.sales = safe_float(_cell(pl, m - 1, c_sales))

    if inp.eps is None:
        raise ValueError(f"PL の最終行（{m + 1} 行目）から EPS を読み取れません。")
    inp.roe_hist = _actual_series(pl, c_roe, scale=0.01)

    # ---- 証券コード・会社名（元スクリプトは PL の A3） ----
    label = None
    if len(pl) > 2 and pl[2] and CODE_RE.search(pl[2][0]):  # 元スクリプトと同じ PL の A3 を優先
        label = pl[2][0].strip()
    for row in ([] if label else pl[:5] + bs[:5]):
        for cell in row:
            if CODE_RE.search(cell or ""):
                label = cell.strip()
                break
        if label:
            break
    if label is None and len(pl) > 2 and pl[2]:
        label = pl[2][0].strip()
    if label:
        mcode = CODE_RE.search(label)
        inp.code = mcode.group(1) if mcode else label[:4]
        inp.company = label
    return inp


# ============================================================
# 外部データ取得
# ============================================================
def normalize_code(code):
    code = (code or "").strip().upper().replace(".T", "")
    code = code.translate(str.maketrans("０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ",
                                        "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    if not re.fullmatch(r"[0-9][0-9A-Z]{2}[0-9A-Z]", code):
        raise ValueError("銘柄コードは 4 桁（例: 7203、130A）で入力してください。")
    return code


def fetch_irbank_csv(code):
    """IR BANK から fy-balance-sheet.csv / fy-profit-and-loss.csv を取得する"""
    out = {}
    for name in ("fy-balance-sheet.csv", "fy-profit-and-loss.csv"):
        url = f"https://f.irbank.net/files/{code}/{name}"
        resp = requests.get(url, headers=HEADERS, timeout=20)
        resp.raise_for_status()
        if b"<html" in resp.content[:500].lower():
            raise ValueError(f"IR BANK から CSV ではなく HTML が返されました: {url}")
        out[name] = resp.content
    return out["fy-balance-sheet.csv"], out["fy-profit-and-loss.csv"]


def fetch_stock_price(code):
    """Yahoo Finance v8 chart API で株価と銘柄名を取得する"""
    url = f"https://query1.finance.yahoo.com/v8/finance/chart/{code}.T?interval=1d&range=1d"
    resp = requests.get(url, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    result = resp.json().get("chart", {}).get("result") or []
    if not result:
        raise ValueError("Yahoo Finance から株価データを取得できませんでした。")
    meta = result[0].get("meta", {})
    price = meta.get("regularMarketPrice")
    if price is None:
        raise ValueError("Yahoo Finance の応答に株価がありません。")
    name = meta.get("longName") or meta.get("shortName") or ""
    return float(price), name


def fetch_yahoo_timeseries(code, types, years=5):
    """Yahoo Finance fundamentals-timeseries から年次データの最新値を取得する"""
    now_ts = int(dt.datetime.now().timestamp())
    url = f"https://query1.finance.yahoo.com/ws/fundamentals-timeseries/v1/finance/timeseries/{code}.T"
    params = {"type": ",".join(types), "period1": now_ts - years * 365 * 86400, "period2": now_ts}
    resp = requests.get(url, params=params, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    latest = {}
    for item in resp.json().get("timeseries", {}).get("result", []) or []:
        for t in types:
            values = [v for v in (item.get(t) or []) if v]
            if values:
                v = max(values, key=lambda x: x.get("asOfDate", ""))
                raw = v.get("reportedValue", {}).get("raw")
                if raw is not None:
                    latest[t] = float(raw)
    return latest


def fetch_current_assets_million(code):
    data = fetch_yahoo_timeseries(code, ["annualCurrentAssets"])
    ca = data.get("annualCurrentAssets")
    return int(round(ca / 1_000_000, 0)) if ca is not None else None


def inputs_from_yahoo(code):
    """IR BANK が使えない場合の代替：Yahoo Finance の年次財務データから同じ指標を組み立てる"""
    types = [
        "annualTotalAssets", "annualStockholdersEquity", "annualTotalEquityGrossMinorityInterest",
        "annualOrdinarySharesNumber", "annualShareIssued", "annualBasicEPS", "annualDilutedEPS",
        "annualNetIncome", "annualNetIncomeCommonStockholders", "annualTotalRevenue",
        "annualCurrentAssets",
    ]
    d = fetch_yahoo_timeseries(code, types)
    inp = FinancialInputs(code=code, source="Yahoo Finance（財務データから算出）")
    total_assets = d.get("annualTotalAssets")
    equity = d.get("annualStockholdersEquity")
    shares = d.get("annualOrdinarySharesNumber") or d.get("annualShareIssued")
    net_income = d.get("annualNetIncomeCommonStockholders") or d.get("annualNetIncome")
    if not (total_assets and equity and shares):
        raise ValueError("Yahoo Finance から総資産・自己資本・発行済株式数を取得できませんでした。")
    inp.total_assets = total_assets
    inp.net_assets = d.get("annualTotalEquityGrossMinorityInterest") or equity
    inp.shareholders_equity = equity
    inp.shares_outstanding = shares
    inp.bps = equity / shares
    inp.equity_ratio = equity / total_assets
    inp.eps = d.get("annualBasicEPS") or d.get("annualDilutedEPS") or (net_income / shares if net_income else None)
    inp.roa = net_income / total_assets if net_income is not None else None
    inp.sales = d.get("annualTotalRevenue")
    ca = d.get("annualCurrentAssets")
    inp.current_assets_million = int(round(ca / 1_000_000, 0)) if ca is not None else None
    inp.notes.append("IR BANK の CSV を取得できなかったため、Yahoo Finance の直近年次実績から算出しました"
                     "（EPS は実績値。IR BANK の予想 EPS とは異なる場合があります）。")
    if inp.eps is None:
        raise ValueError("Yahoo Finance から EPS / 純利益を取得できませんでした。")
    return inp


def _actual_series(rows, col, scale=1.0):
    """「年度」見出しより下の実績行（予想行を除く）から 1 列の数値を古い順に取り出す"""
    h = next((i for i, row in enumerate(rows[:15]) if row and row[0].strip() == "年度"), 0)
    out = []
    for row in rows[h + 1:]:
        if not row or not re.search(r"\d{4}", row[0]) or any("予" in v for v in row):
            continue
        v = safe_float(row[col]) if col < len(row) else None
        if v is not None:
            out.append(v * scale)
    return out


def parse_irbank_dividend(data):
    """IR BANK の fy-stock-dividend.csv から直近 3 期の実績配当性向の中央値（小数）を返す。無ければ None"""
    rows = read_csv_rows(data)
    col = _header_index(rows, ("配当性向",), None)
    if col is None:
        return None
    vals = [v for v in _actual_series(rows, col, scale=0.01) if 0 <= v <= 1.5][-3:]
    return min(sorted(vals)[len(vals) // 2], 1.0) if vals else None


def fetch_irbank_dividend(code):
    url = f"https://f.irbank.net/files/{code}/fy-stock-dividend.csv"
    resp = requests.get(url, headers=HEADERS, timeout=20)
    resp.raise_for_status()
    if b"<html" in resp.content[:500].lower():
        raise ValueError("配当の CSV ではなく HTML が返されました。")
    return parse_irbank_dividend(resp.content)


def load_inputs_by_code(code):
    """銘柄コードから入力データを取得（IR BANK → 失敗時 Yahoo Finance）"""
    code = normalize_code(code)
    errors = []
    try:
        bs, pl = fetch_irbank_csv(code)
        inp = parse_irbank_csv(bs, pl)
        inp.code = code
        inp.source = "IR BANK（CSV）"
        inp.raw_pl = pl
        try:
            inp.payout_ratio = fetch_irbank_dividend(code)
        except Exception:  # noqa: BLE001
            pass
        if inp.payout_ratio is None:
            inp.notes.append(f"配当性向を取得できなかったため {DEFAULT_PAYOUT * 100:.0f}% と仮定しました（画面で変更できます）。")
        return inp, errors
    except Exception as e:  # noqa: BLE001
        errors.append(f"IR BANK からの取得に失敗: {e}")
    inp = inputs_from_yahoo(code)
    return inp, errors


def enrich_with_market_data(inp):
    """株価・流動資産を Yahoo Finance から補完する。取得できなかった項目は警告として返す"""
    warnings = []
    if inp.stock_price is None and inp.code:
        try:
            price, name = fetch_stock_price(inp.code)
            inp.stock_price = price
            if not inp.company:
                inp.company = f"{inp.code} {name}".strip()
        except Exception as e:  # noqa: BLE001
            warnings.append(f"株価を取得できませんでした（{e}）。画面で株価を入力してください。")
    if inp.current_assets_million is None and inp.code:
        try:
            inp.current_assets_million = fetch_current_assets_million(inp.code)
            if inp.current_assets_million is None:
                warnings.append("流動資産を取得できませんでした。画面で入力すると「時価総額/正味価値」を計算します。")
        except Exception as e:  # noqa: BLE001
            warnings.append(f"流動資産を取得できませんでした（{e}）。画面で入力すると「時価総額/正味価値」を計算します。")
    if not inp.company:
        inp.company = inp.code
    return warnings


# ============================================================
# 計算：残余利益モデル（Residual Income Model / Edwards-Bell-Ohlson）
#
#   理論株価 V0 = B0 + Σ_{t=1..T} (ROE_t − r) × B_{t−1} / (1 + r)^t
#
#   B0    : 直近の BPS（1 株当たり純資産）
#   ROE_1 : 予想 EPS ÷ BPS（または過去 3 期平均 ROE）
#   ROE_t : T 年かけて株主資本コスト r まで直線的に収れん（超過利益は競争で消えていく、という保守的な前提）
#   B_t   : B_{t−1} × (1 + ROE_t × (1 − 配当性向))（クリーンサープラス関係）
#   T 年目に ROE = r となるので、それ以降の残余利益（継続価値）は 0
# ============================================================
ROE_BASIS_FORECAST = "予想（直近）EPS ÷ BPS"
ROE_BASIS_AVG3 = "過去3期平均 ROE"
ROE_BASES = (ROE_BASIS_FORECAST, ROE_BASIS_AVG3)

SENS_COE = (0.06, 0.07, 0.08, 0.09, 0.10)
SENS_YEARS = (5, 10, 15)

DISCLAIMER = (
    "本ツールの理論株価は、公開されている財務データを一般的な評価手法（残余利益モデル）と一定の前提条件で"
    "機械的に計算した参考値です。特定の銘柄の売買を勧誘・推奨するものではなく、将来の株価や運用成果を保証するものでもありません。"
    "データの正確性・完全性は保証しません。投資に関する最終決定は、ご自身の判断と責任で行ってください。"
)

METHOD_SOURCES = [
    "Ohlson, J. A. (1995) “Earnings, Book Values, and Dividends in Equity Valuation,” Contemporary Accounting Research 11(2).",
    "Frankel, R. & Lee, C. M. C. (1998) “Accounting Valuation, Market Expectation, and Cross-sectional Stock Returns,” Journal of Accounting and Economics 25(3).",
    "Lee, C. M. C., Myers, J. & Swaminathan, B. (1999) “What is the Intrinsic Value of the Dow?,” Journal of Finance 54(5).",
    "Francis, J., Olsson, P. & Oswald, D. (2000) “Comparing the Accuracy and Explainability of Dividend, Free Cash Flow, and Abnormal Earnings Equity Value Estimates,” Journal of Accounting Research 38(1).",
    "日本公認会計士協会『企業価値評価ガイドライン』（経営研究調査会研究報告第32号、2013年改正）：インカム・アプローチの一つとして残余利益法を掲載。",
    "経済産業省『持続的成長への競争力とインセンティブ ～企業と投資家の望ましい関係構築～』（伊藤レポート、2014年）：グローバルな投資家が想定する資本コストを踏まえ ROE 8% を最低ラインとして提示。",
]


def residual_income_value(bps, roe1, coe=COST_OF_EQUITY, years=FADE_YEARS, payout=DEFAULT_PAYOUT):
    """残余利益モデルで 1 株当たり理論株価を計算する。戻り値: (理論株価, 年度別の内訳リスト)"""
    if bps is None or roe1 is None or bps <= 0:
        return None, []
    years = max(int(years), 1)
    payout = min(max(payout, 0.0), 1.0)
    rows, book, total = [], bps, 0.0
    for t in range(1, years + 1):
        roe_t = roe1 + (coe - roe1) * (t - 1) / (years - 1) if years > 1 else roe1
        ri = (roe_t - coe) * book
        df = 1 / (1 + coe) ** t
        rows.append({"t": t, "roe": roe_t, "book": book, "ri": ri, "df": df, "pv": ri * df})
        total += ri * df
        book = book * (1 + roe_t * (1 - payout))
    return bps + total, rows


@dataclass
class Result:
    # 前提
    roe_basis: str = ROE_BASIS_FORECAST
    roe_used: float = None            # ROE_1（小数）
    cost_of_equity: float = COST_OF_EQUITY
    years: int = FADE_YEARS
    payout: float = DEFAULT_PAYOUT
    # 理論株価
    book_value: float = None          # B0（BPS）
    pv_residual_income: float = None  # 残余利益の現在価値合計
    theoretical_price: float = None
    ri_rows: list = field(default_factory=list)
    price_low: float = None           # r + 1%
    price_high: float = None          # r − 1%
    sensitivity: dict = field(default_factory=dict)  # {(r, T): 理論株価}
    # 株価との比較
    stock_price: float = None
    upside_pct: float = None          # 理論株価 ÷ 株価 − 1（%）
    gap_pct: float = None             # 株価 ÷ 理論株価 − 1（%）。−30% 以下で割安圏、+30% 以上で割高圏
    verdict: str = ""
    recommend: bool = None            # 割安圏なら True、割高圏なら False、それ以外 None（画面の色分け用）
    summary: str = ""
    # 参考指標
    market_cap_million: float = None
    sales_million: float = None
    pbr: float = None
    per: float = None
    psr: float = None
    liabilities_million: float = None
    current_assets_million: float = None
    net_value_million: float = None   # 正味流動資産（流動資産 − 負債）
    mc_to_net_value: float = None
    upside_ok: bool = None
    mc_net_ok: bool = None
    notes: list = field(default_factory=list)


def select_roe(inp, basis):
    """ROE_1 を選ぶ。戻り値: (ROE, 実際に使った基準)"""
    fwd = inp.eps / inp.bps if inp.eps is not None and inp.bps else None
    avg3 = sum(inp.roe_hist[-3:]) / len(inp.roe_hist[-3:]) if inp.roe_hist else None
    if basis == ROE_BASIS_AVG3 and avg3 is not None:
        return avg3, ROE_BASIS_AVG3
    if fwd is not None:
        return fwd, ROE_BASIS_FORECAST
    return avg3, ROE_BASIS_AVG3 if avg3 is not None else basis


def calculate(inp, coe=COST_OF_EQUITY, years=FADE_YEARS, payout=None, roe_basis=ROE_BASIS_FORECAST):
    if payout is None:
        payout = inp.payout_ratio if inp.payout_ratio is not None else DEFAULT_PAYOUT
    r = Result(cost_of_equity=coe, years=int(years), payout=payout)
    r.roe_used, r.roe_basis = select_roe(inp, roe_basis)
    if roe_basis == ROE_BASIS_AVG3 and r.roe_basis != ROE_BASIS_AVG3:
        r.notes.append("過去の ROE を読み取れなかったため、予想（直近）EPS ÷ BPS を使いました。")
    r.book_value = inp.bps

    if inp.bps is not None and inp.bps <= 0:
        r.notes.append("BPS がマイナス（債務超過）のため、残余利益モデルでは理論株価を計算できません。")
    elif r.roe_used is None:
        r.notes.append("EPS・ROE を読み取れないため、理論株価を計算できません。")
    else:
        r.theoretical_price, r.ri_rows = residual_income_value(inp.bps, r.roe_used, coe, years, payout)
        r.pv_residual_income = r.theoretical_price - inp.bps
        r.price_low, _ = residual_income_value(inp.bps, r.roe_used, coe + 0.01, years, payout)
        r.price_high, _ = residual_income_value(inp.bps, r.roe_used, max(coe - 0.01, 0.01), years, payout)
        r.sensitivity = {(c, y): residual_income_value(inp.bps, r.roe_used, c, y, payout)[0]
                         for c in SENS_COE for y in SENS_YEARS}
        if r.theoretical_price <= 0:
            r.notes.append("赤字が大きく、理論株価がマイナスになりました。収益力からは株主価値を説明できない状態です。")

    r.stock_price = inp.stock_price
    tp = r.theoretical_price
    if r.stock_price and tp and tp > 0:
        r.upside_pct = (tp / r.stock_price - 1) * 100
        r.gap_pct = (r.stock_price / tp - 1) * 100
        r.upside_ok = r.gap_pct <= -UPSIDE_THRESHOLD
        if r.gap_pct <= -UPSIDE_THRESHOLD:
            r.verdict, r.recommend = "割安圏", True
        elif r.gap_pct >= UPSIDE_THRESHOLD:
            r.verdict, r.recommend = "割高圏", False
        else:
            r.verdict = "概ね理論株価の範囲内"
        r.summary = (f"株価 {r.stock_price:,.0f} 円は理論株価 {tp:,.0f} 円"
                     f"（株主資本コスト {coe * 100:.1f}%±1% で {r.price_low:,.0f}〜{r.price_high:,.0f} 円）に対して "
                     f"{r.gap_pct:+.1f}% の水準で、評価は「{r.verdict}」です"
                     f"（±{UPSIDE_THRESHOLD}% を超える乖離を割安圏・割高圏としています）。")
    elif tp is not None and tp <= 0:
        r.verdict, r.recommend = "評価不能", False
        r.summary = "理論株価がプラスにならないため、株価との比較はできません。"
    else:
        r.verdict = "判断保留"
        r.summary = "株価または理論株価が無いため、比較は保留します。株価を入力してください。"

    # ---- 参考指標（PBR・PER・PSR・正味流動資産） ----
    shares = inp.shares_outstanding
    if shares is None and inp.shareholders_equity and inp.bps:
        shares = inp.shareholders_equity / inp.bps
        inp.shares_outstanding = shares
    if r.stock_price and shares:
        r.market_cap_million = int(round(r.stock_price * shares / 1_000_000, 0))
    if r.stock_price and inp.bps and inp.bps > 0:
        r.pbr = r.stock_price / inp.bps
    if r.stock_price and inp.eps and inp.eps > 0:
        r.per = r.stock_price / inp.eps
    if inp.sales is not None:
        r.sales_million = int(round(inp.sales / 1_000_000, 0))
    if r.market_cap_million and r.sales_million:
        r.psr = r.market_cap_million / r.sales_million
    if inp.total_assets is not None and inp.net_assets is not None:
        r.liabilities_million = int(round((inp.total_assets - inp.net_assets) / 1_000_000, 0))
    r.current_assets_million = inp.current_assets_million
    if r.current_assets_million is not None and r.liabilities_million is not None:
        r.net_value_million = r.current_assets_million - r.liabilities_million
        if r.net_value_million != 0 and r.market_cap_million:
            r.mc_to_net_value = r.market_cap_million / r.net_value_million
            r.mc_net_ok = 0 < r.mc_to_net_value < 1
    return r


# ============================================================
# Excel 出力（計算式をセルに残し、誰でも検算できる形にする）
# ============================================================
def build_excel(inp, r):
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = px.Workbook()
    ws = wb.active
    ws.title = "理論株価"
    bold, title = Font(bold=True), Font(bold=True, size=14)
    head_fill = PatternFill(patternType="solid", fgColor="DDEBF7")
    input_fill = PatternFill(patternType="solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = f"理論株価（残余利益モデル）：{inp.company or inp.code}"
    ws["A1"].font = title
    ws["A2"] = f"データ元：{inp.source}　作成日：{dt.date.today():%Y/%m/%d}　黄色のセルは前提条件で、書き換えると再計算されます。"

    # ---- 前提条件 ----
    ws["A4"] = "前提条件"
    ws["A4"].font = bold
    inputs = [
        ("BPS（1株当たり純資産・直近実績, 円）", inp.bps, "#,##0.00"),
        ("ROE（1年目）", r.roe_used, "0.00%"),
        ("株主資本コスト r", r.cost_of_equity, "0.0%"),
        ("ROE が r に収れんするまでの年数 T", r.years, "0"),
        ("配当性向", r.payout, "0.0%"),
        ("現在の株価（円）", r.stock_price, "#,##0"),
    ]
    for i, (k, v, f) in enumerate(inputs):
        row = 5 + i
        ws.cell(row=row, column=1, value=k)
        c = ws.cell(row=row, column=3, value=v)
        c.number_format, c.fill = f, input_fill
    ws["D6"] = f"基準：{r.roe_basis}" + (f"（EPS {inp.eps:,.2f} 円 ÷ BPS）" if r.roe_basis == ROE_BASIS_FORECAST and inp.eps is not None else "")
    # C5 BPS, C6 ROE1, C7 r, C8 T, C9 配当性向, C10 株価

    # ---- 年度別の残余利益 ----
    top = 12
    for c, h in enumerate(["年", "ROE", "期首 BPS", "残余利益 (ROE−r)×BPS", "割引係数 1/(1+r)^t", "現在価値"], 1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font, cell.fill = bold, head_fill
    n = len(r.ri_rows)
    for i in range(n):
        row = top + 1 + i
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=f"=$C$6+($C$7-$C$6)*(A{row}-1)/MAX($C$8-1,1)").number_format = "0.00%"
        ws.cell(row=row, column=3, value="=$C$5" if i == 0 else f"=C{row - 1}*(1+B{row - 1}*(1-$C$9))").number_format = "#,##0.00"
        ws.cell(row=row, column=4, value=f"=(B{row}-$C$7)*C{row}").number_format = "#,##0.00"
        ws.cell(row=row, column=5, value=f"=1/(1+$C$7)^A{row}").number_format = "0.0000"
        ws.cell(row=row, column=6, value=f"=D{row}*E{row}").number_format = "#,##0.00"
    last = top + max(n, 1)

    s = last + 2
    ws.cell(row=s, column=1, value="残余利益の現在価値の合計（円）").font = bold
    ws.cell(row=s, column=3, value=f"=SUM(F{top + 1}:F{last})" if n else None).number_format = "#,##0.00"
    ws.cell(row=s + 1, column=1, value="理論株価 = BPS + 残余利益の現在価値（円）").font = bold
    tp = ws.cell(row=s + 1, column=3, value=f"=C5+C{s}" if n else None)
    tp.number_format, tp.font = "#,##0", Font(bold=True, size=12)
    ws.cell(row=s + 2, column=1, value="株価の理論株価からの乖離（株価 ÷ 理論株価 − 1）").font = bold
    gap = ws.cell(row=s + 2, column=3, value=f'=IF(AND(C10>0,C{s + 1}>0),C10/C{s + 1}-1,"-")' if n else "-")
    gap.number_format = "+0.0%;-0.0%"
    ws.cell(row=s + 3, column=1, value="評価").font = bold
    th = UPSIDE_THRESHOLD / 100
    ws.cell(row=s + 3, column=3, value=(
        f'=IF(ISNUMBER(C{s + 2}),IF(C{s + 2}<=-{th},"割安圏",IF(C{s + 2}>={th},"割高圏","概ね理論株価の範囲内")),"-")'
        if n else r.verdict))
    ws.cell(row=s + 3, column=4, value=f"乖離が ±{UPSIDE_THRESHOLD}% を超えたら割安圏・割高圏（安全余裕率の考え方）")
    if r.theoretical_price is not None:
        ws.cell(row=s + 4, column=1, value="（参考）Python で計算した理論株価")
        ws.cell(row=s + 4, column=3, value=round(r.theoretical_price, 2)).number_format = "#,##0"

    # ---- 感応度 ----
    t = s + 6
    ws.cell(row=t, column=1, value="感応度：株主資本コスト r と収れん年数 T を変えたときの理論株価（円）").font = bold
    ws.cell(row=t + 1, column=1, value="r ＼ T").font = bold
    for j, y in enumerate(SENS_YEARS):
        cell = ws.cell(row=t + 1, column=2 + j, value=f"{y} 年")
        cell.font, cell.fill = bold, head_fill
    for i, c in enumerate(SENS_COE):
        ws.cell(row=t + 2 + i, column=1, value=c).number_format = "0.0%"
        for j, y in enumerate(SENS_YEARS):
            v = r.sensitivity.get((c, y))
            ws.cell(row=t + 2 + i, column=2 + j, value=round(v) if v is not None else "-").number_format = "#,##0"

    # ---- 参考指標 ----
    q = t + 3 + len(SENS_COE)
    ws.cell(row=q, column=1, value="参考指標").font = bold
    ref = [
        ("PBR（株価 ÷ BPS, 倍）", r.pbr, "0.00"),
        ("PER（株価 ÷ EPS, 倍）", r.per, "0.0"),
        ("PSR（時価総額 ÷ 売上高, 倍）", r.psr, "0.00"),
        ("自己資本比率", inp.equity_ratio, "0.0%"),
        ("時価総額（百万円）", r.market_cap_million, "#,##0"),
        ("売上高（百万円）", r.sales_million, "#,##0"),
        ("流動資産（百万円）", r.current_assets_million, "#,##0"),
        ("負債（総資産 − 純資産, 百万円）", r.liabilities_million, "#,##0"),
        ("時価総額 ÷ 正味流動資産（倍）", r.mc_to_net_value, "0.00"),
    ]
    for i, (k, v, f) in enumerate(ref):
        ws.cell(row=q + 1 + i, column=1, value=k)
        c = ws.cell(row=q + 1 + i, column=3, value=v if v is not None else "-")
        if v is not None:
            c.number_format = f

    # ---- 注記・出典・免責 ----
    z = q + len(ref) + 2
    notes = [
        ("算定式", "理論株価 = BPS + Σ (ROE_t − r) × 期首BPS_t ÷ (1 + r)^t。ROE は T 年かけて r まで直線的に低下し、"
                 "BPS は利益から配当を除いた分だけ増える（クリーンサープラス）と仮定。T 年目以降の超過利益は 0。"),
        ("前提の根拠", f"株主資本コスト {COST_OF_EQUITY * 100:.0f}% は伊藤レポート（経済産業省, 2014）が示した ROE の最低ライン。"
                   f"配当性向はデータから取得できない場合 {DEFAULT_PAYOUT * 100:.0f}% と仮定。"),
    ] + [("出典", s_) for s_ in METHOD_SOURCES] + [("免責事項", DISCLAIMER)]
    for i, (k, v) in enumerate(notes):
        ws.cell(row=z + i, column=1, value=k).font = bold
        c = ws.cell(row=z + i, column=2, value=v)
        c.alignment = wrap
        ws.merge_cells(start_row=z + i, start_column=2, end_row=z + i, end_column=6)
        ws.row_dimensions[z + i].height = 45

    ws.column_dimensions["A"].width = 42
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_file_name(inp):
    name = (inp.company or inp.code or "result").strip()
    name = re.sub(r'[\\/:*?"<>|]', "_", name)
    return f"StockPrice_{name}.xlsx"
