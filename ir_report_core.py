# -*- coding: utf-8 -*-
"""
ir_report_core.py
IR_report.ipynb（Google Colab）の処理を Web アプリから呼べるよう関数化したもの。

- IR BANK の fy-profit-and-loss.csv（損益計算書の年次推移）の読み込み（'-' や空欄は欠損として扱う）
- 増収率・増益率・成長性判定、直近3期平均、EPS 年平均成長率
- 各年度の期末株価を Yahoo Finance の月足から取得（Close.xlsx は不要）
- EPS と株価の相関、過去 PER の水準から見た割安/割高の目安
- Excel 出力（表・売上高/営業利益グラフ・EPS と株価の散布図）
"""

import datetime as dt
import io
import re
import statistics
from dataclasses import dataclass, field

import openpyxl as px
import requests
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from stockprice_core import HEADERS, normalize_code, read_csv_rows, safe_float

# 成長性判定の基準（直近3期平均の 増収率 + 増益率）
GROWTH_HIGH = 20.0  # % 以上で「高成長」

FY_RE = re.compile(r"(\d{4})\D+(\d{1,2})")


# ============================================================
# データ構造
# ============================================================
@dataclass
class YearRow:
    label: str                   # 年度（例 "2024/03"）
    year: int = None
    month: int = None
    forecast: bool = False       # 予想行
    sales: float = None          # 売上高（円）
    op_income: float = None      # 営業利益（円）
    net_income: float = None     # 純利益（円）
    eps: float = None            # EPS（円）
    sales_growth: float = None   # 増収率（%）
    op_growth: float = None      # 増益率（営業利益、%）
    growth_score: float = None   # 成長性 = 増収率 + 増益率（%）
    pattern: str = ""            # 増収増益 など
    price: float = None          # 期末株価（円）
    per: float = None            # 期末 PER（倍）


@dataclass
class IRReport:
    code: str = ""
    company: str = ""
    source: str = ""
    rows: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    # まとめ
    avg_sales_growth_3y: float = None
    avg_op_growth_3y: float = None
    growth_score_3y: float = None
    growth_label: str = ""
    eps_cagr: float = None
    eps_cagr_years: int = None
    eps_price_corr: float = None
    per_median: float = None
    per_min: float = None
    per_max: float = None
    current_price: float = None
    current_eps: float = None
    current_eps_label: str = ""
    current_per: float = None
    per_fair_price: float = None   # 過去 PER 中央値 × 直近 EPS
    per_upside_pct: float = None

    @property
    def actual_rows(self):
        return [r for r in self.rows if not r.forecast]


# ============================================================
# 読み込み
# ============================================================
def _find_col(header, names):
    for i, h in enumerate(header):
        if h.strip() in names:
            return i
    return None


def parse_pl_history(pl_data):
    """IR BANK の fy-profit-and-loss.csv を年度ごとの行に変換する。'-' や空欄は欠損（None）にする"""
    rows = read_csv_rows(pl_data)
    h = next((i for i, r in enumerate(rows[:15]) if r and r[0].strip() == "年度"), None)
    if h is None:
        raise ValueError("損益計算書 CSV に「年度」の見出し行が見つかりません。")
    header = rows[h]
    c_sales = _find_col(header, ("売上高", "営業収益", "経常収益", "売上収益"))
    c_op = _find_col(header, ("営業利益", "経常利益"))
    c_net = _find_col(header, ("純利益", "当期純利益"))
    c_eps = _find_col(header, ("EPS",))
    if c_sales is None:
        raise ValueError("損益計算書 CSV に「売上高」の列が見つかりません。")

    rep = IRReport(source="IR BANK（CSV）")
    for r in rows[h + 1:]:
        if not r or not r[0].strip():
            continue
        label = r[0].strip()
        m = FY_RE.search(label)
        if not m:
            continue
        get = lambda c: safe_float(r[c]) if c is not None and c < len(r) else None  # noqa: E731
        rep.rows.append(YearRow(
            label=label,
            year=int(m.group(1)), month=int(m.group(2)),
            forecast=any("予" in v for v in r),
            sales=get(c_sales), op_income=get(c_op), net_income=get(c_net), eps=get(c_eps),
        ))
    if not rep.rows:
        raise ValueError("損益計算書 CSV に年度データがありません。")

    # 銘柄コード・会社名（元ノートブックは A3）
    for r in rows[:h]:
        for cell in r:
            m = re.search(r"(?<![0-9A-Za-z/])([0-9][0-9A-Z]{2}[0-9A-Z])(?![0-9A-Za-z/])", cell or "")
            if m:
                rep.code, rep.company = m.group(1), cell.strip()
                break
        if rep.code:
            break
    return rep


def _fetch_yahoo_series(code, types, years=12):
    now_ts = int(dt.datetime.now().timestamp())
    url = f"https://query1.finance.yahoo.com/ws/fundamentals-timeseries/v1/finance/timeseries/{code}.T"
    params = {"type": ",".join(types), "period1": now_ts - years * 366 * 86400, "period2": now_ts}
    resp = requests.get(url, params=params, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    out = {}
    for item in resp.json().get("timeseries", {}).get("result", []) or []:
        for t in types:
            for v in item.get(t) or []:
                if v and v.get("asOfDate") and v.get("reportedValue", {}).get("raw") is not None:
                    out.setdefault(v["asOfDate"][:7], {})[t] = float(v["reportedValue"]["raw"])
    return out


def pl_history_from_yahoo(code):
    """IR BANK が使えない場合の代替：Yahoo Finance の年次 P/L（通常 4〜5 期分）"""
    t = ["annualTotalRevenue", "annualOperatingIncome", "annualNetIncome", "annualBasicEPS", "annualDilutedEPS"]
    d = _fetch_yahoo_series(code, t)
    if not d:
        raise ValueError("Yahoo Finance から損益データを取得できませんでした。")
    rep = IRReport(code=code, source="Yahoo Finance（年次 P/L）")
    for ym in sorted(d):
        v = d[ym]
        y, mo = ym.split("-")
        rep.rows.append(YearRow(
            label=f"{y}/{mo}", year=int(y), month=int(mo),
            sales=v.get("annualTotalRevenue"), op_income=v.get("annualOperatingIncome"),
            net_income=v.get("annualNetIncome"), eps=v.get("annualBasicEPS") or v.get("annualDilutedEPS"),
        ))
    rep.notes.append("IR BANK の CSV を取得できなかったため、Yahoo Finance の年次データ（直近数期分のみ）を使いました。")
    return rep


def fetch_month_end_prices(code, years=20):
    """Yahoo Finance の月足から {(年, 月): 月末終値} と現在値を返す"""
    url = f"https://query1.finance.yahoo.com/v8/finance/chart/{code}.T?interval=1mo&range={years}y"
    resp = requests.get(url, headers=HEADERS, timeout=15)
    resp.raise_for_status()
    result = (resp.json().get("chart", {}).get("result") or [None])[0]
    if not result:
        raise ValueError("Yahoo Finance から株価の履歴を取得できませんでした。")
    ts = result.get("timestamp") or []
    closes = (result.get("indicators", {}).get("adjclose") or [{}])[0].get("adjclose") \
        or (result.get("indicators", {}).get("quote") or [{}])[0].get("close") or []
    jst = dt.timezone(dt.timedelta(hours=9))
    prices = {}
    for t, c in zip(ts, closes):
        if c is not None:
            d = dt.datetime.fromtimestamp(t, jst)
            prices[(d.year, d.month)] = float(c)
    current = result.get("meta", {}).get("regularMarketPrice")
    return prices, (float(current) if current is not None else None)


# ============================================================
# 分析
# ============================================================
def _pct(cur, prev):
    """前期比（%）。前期がゼロ・マイナスのときは率が意味を持たないので None"""
    if cur is None or prev is None or prev <= 0:
        return None
    return (cur / prev - 1) * 100


def _pattern(sg, og):
    if sg is None or og is None:
        return "-"
    return ("増収" if sg >= 0 else "減収") + ("増益" if og >= 0 else "減益")


def _avg(vals):
    vals = [v for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


def _corr(xs, ys):
    pairs = [(x, y) for x, y in zip(xs, ys) if x is not None and y is not None]
    if len(pairs) < 3:
        return None
    # statistics.correlation は Python 3.10 以降のみなので自前で計算する
    n = len(pairs)
    mx, my = sum(p[0] for p in pairs) / n, sum(p[1] for p in pairs) / n
    sxy = sum((x - mx) * (y - my) for x, y in pairs)
    sxx = sum((x - mx) ** 2 for x, _ in pairs)
    syy = sum((y - my) ** 2 for _, y in pairs)
    if sxx == 0 or syy == 0:
        return None
    return sxy / (sxx * syy) ** 0.5


def analyze(rep, prices=None, current_price=None):
    """増収率・増益率・成長性判定、期末株価・PER、EPS と株価の相関などを計算する"""
    prices = prices or {}
    prev = None
    for r in rep.rows:
        if prev is not None:
            r.sales_growth = _pct(r.sales, prev.sales)
            r.op_growth = _pct(r.op_income, prev.op_income)
            if r.sales_growth is not None and r.op_growth is not None:
                r.growth_score = r.sales_growth + r.op_growth
            r.pattern = _pattern(r.sales_growth, r.op_growth)
        if not r.forecast and r.year:
            r.price = prices.get((r.year, r.month))
            if r.price and r.eps and r.eps > 0:
                r.per = r.price / r.eps
        prev = r

    actual = rep.actual_rows
    last3 = actual[-3:]
    rep.avg_sales_growth_3y = _avg([r.sales_growth for r in last3])
    rep.avg_op_growth_3y = _avg([r.op_growth for r in last3])
    if rep.avg_sales_growth_3y is not None and rep.avg_op_growth_3y is not None:
        rep.growth_score_3y = rep.avg_sales_growth_3y + rep.avg_op_growth_3y
        if rep.growth_score_3y >= GROWTH_HIGH and rep.avg_sales_growth_3y > 0 and rep.avg_op_growth_3y > 0:
            rep.growth_label = "高成長"
        elif rep.avg_sales_growth_3y >= 0 and rep.avg_op_growth_3y >= 0:
            rep.growth_label = "安定成長"
        elif rep.growth_score_3y >= 0:
            rep.growth_label = "成長鈍化"
        else:
            rep.growth_label = "業績悪化"
    else:
        rep.growth_label = "判定不可"

    # EPS 年平均成長率（最大10年。始点・終点ともプラスのときのみ）
    eps_rows = [r for r in actual if r.eps is not None]
    if len(eps_rows) >= 2:
        end = eps_rows[-1]
        start = next((r for r in eps_rows[-11:] if r.eps > 0), None)
        if start is not None and start is not end and end.eps > 0:
            n = end.year - start.year + (end.month - start.month) / 12
            if n > 0:
                rep.eps_cagr = ((end.eps / start.eps) ** (1 / n) - 1) * 100
                rep.eps_cagr_years = round(n)

    rep.eps_price_corr = _corr([r.eps for r in actual], [r.price for r in actual])
    pers = [r.per for r in actual if r.per is not None]
    if pers:
        rep.per_median = statistics.median(pers)
        rep.per_min, rep.per_max = min(pers), max(pers)
    if actual and not any(r.price for r in actual):
        rep.notes.append("期末株価を取得できなかったため、株価・PER・相関は空欄です。")

    eps_src = next((r for r in reversed(rep.rows) if r.forecast and r.eps is not None), None) \
        or (eps_rows[-1] if eps_rows else None)
    if eps_src is not None:
        rep.current_eps = eps_src.eps
        rep.current_eps_label = f"{eps_src.label}{'（予想）' if eps_src.forecast and '予' not in eps_src.label else ''}"
    set_current_price(rep, current_price)
    return rep


def set_current_price(rep, price):
    """現在の株価から現在 PER と、過去 PER 中央値で見た目安株価・上昇余地を計算する（予想 EPS があれば優先）"""
    rep.current_price = price or None
    rep.current_per = rep.per_fair_price = rep.per_upside_pct = None
    eps = rep.current_eps
    if eps is None or eps <= 0:
        return rep
    if rep.current_price:
        rep.current_per = rep.current_price / eps
    if rep.per_median:
        rep.per_fair_price = rep.per_median * eps
        if rep.current_price:
            rep.per_upside_pct = (rep.per_fair_price / rep.current_price - 1) * 100
    return rep


def load_report(code=None, pl_data=None, with_prices=True):
    """銘柄コード（または PL CSV のバイト列）から IR レポートを作る。戻り値: (report, 警告リスト)"""
    warnings = []
    if pl_data is not None:
        rep = parse_pl_history(pl_data)
    else:
        rep = pl_history_from_yahoo(normalize_code(code))
    if code:
        rep.code = normalize_code(code)
    prices, current = {}, None
    if with_prices and rep.code:
        try:
            prices, current = fetch_month_end_prices(rep.code)
        except Exception as e:  # noqa: BLE001
            warnings.append(f"株価の履歴を取得できませんでした（{e}）。期末株価・PER は空欄になります。")
    return analyze(rep, prices, current), warnings


# ============================================================
# 投資判断の助言（理論株価 × 成長性 × PER 水準）
# ============================================================
def combined_advice(theory_upside_pct, rep, upside_threshold=30):
    """理論株価の上昇余地、成長性、過去 PER との比較からコメントを組み立てる。
    戻り値: (総合評価ラベル, 理由のリスト)"""
    reasons, score, n = [], 0, 0
    if theory_upside_pct is not None:
        n += 1
        if theory_upside_pct >= upside_threshold:
            score += 1
            reasons.append(f"理論株価に対して上昇余地 {theory_upside_pct:.1f}% と十分に割安です（基準 {upside_threshold}%）。")
        elif theory_upside_pct >= 0:
            reasons.append(f"理論株価は現在値を上回るものの、上昇余地 {theory_upside_pct:.1f}% は基準 {upside_threshold}% に届きません。")
        else:
            score -= 1
            reasons.append(f"現在の株価は理論株価を {-theory_upside_pct:.1f}% 上回っており、資産・事業価値から見ると割高です。")
    if rep is not None and rep.growth_label not in ("", "判定不可"):
        n += 1
        g = f"直近3期平均で増収率 {rep.avg_sales_growth_3y:.1f}%・増益率 {rep.avg_op_growth_3y:.1f}% の「{rep.growth_label}」です。"
        if rep.growth_label == "高成長":
            score += 1
        elif rep.growth_label == "業績悪化":
            score -= 1
            g += " 理論株価の前提となる利益が今後も続くかに注意が必要です。"
        reasons.append(g)
    if rep is not None and rep.per_upside_pct is not None:
        n += 1
        p = (f"現在 PER {rep.current_per:.1f} 倍は過去の期末 PER 中央値 {rep.per_median:.1f} 倍"
             f"（{rep.per_min:.1f}〜{rep.per_max:.1f} 倍）と比べて")
        if rep.per_upside_pct >= 10:
            score += 1
            reasons.append(p + "低く、過去の評価水準から見ても割安です。")
        elif rep.per_upside_pct <= -10:
            score -= 1
            reasons.append(p + "高く、過去の評価水準から見ると割高です。")
        else:
            reasons.append(p + "ほぼ同水準です。")
    if rep is not None and rep.eps_price_corr is not None and rep.eps_price_corr < 0.3:
        reasons.append(f"EPS と株価の相関が {rep.eps_price_corr:.2f} と弱く、業績が株価に反映されにくい銘柄です。PER による目安は参考程度にしてください。")

    if n == 0:
        return "判断保留", ["必要なデータが不足しているため、総合評価はできません。"]
    if score >= 2:
        label = "割安・買い検討"
    elif score >= 1:
        label = "やや割安・押し目待ち"
    elif score <= -2:
        label = "割高・見送り"
    elif score <= -1:
        label = "やや割高・様子見"
    else:
        label = "適正水準・様子見"
    return label, reasons


# ============================================================
# Excel 出力
# ============================================================
HEADER = ["年度", "売上高", "営業利益", "増収率", "増益率", "成長性(増収率+増益率)", "増収増益",
          "純利益", "EPS", "期末株価", "PER"]


def write_ir_sheet(ws, rep):
    """ワークシートに IR レポート（表とグラフ）を書き込む"""
    ws["A1"] = f"IR レポート：{rep.company or rep.code}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"データ元：{rep.source}　作成日：{dt.date.today():%Y/%m/%d}　（金額は百万円、率は前期比）"

    top = 4
    for c, h in enumerate(HEADER, 1):
        ws.cell(row=top, column=c, value=h)
    m = lambda v: None if v is None else round(v / 1_000_000)  # noqa: E731
    p = lambda v: None if v is None else v / 100  # noqa: E731
    forecast_fill = PatternFill(patternType="solid", fgColor="FFF2CC")
    for i, r in enumerate(rep.rows):
        row = top + 1 + i
        vals = [r.label + ("（予想）" if r.forecast and "予" not in r.label else ""), m(r.sales), m(r.op_income),
                p(r.sales_growth), p(r.op_growth), p(r.growth_score), r.pattern or "-",
                m(r.net_income), r.eps, r.price, r.per]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=row, column=c, value=v)
            if c in (2, 3, 8):
                cell.number_format = "#,##0"
            elif c in (4, 5, 6):
                cell.number_format = "0.0%"
            elif c in (9, 10, 11):
                cell.number_format = "#,##0.0"
            if r.forecast:
                cell.fill = forecast_fill
    last = top + len(rep.rows)
    table = Table(displayName=f"IR_{re.sub(r'[^0-9A-Za-z]', '', rep.code) or 'report'}", ref=f"A{top}:K{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)

    # まとめ
    s = last + 2
    bold, yellow = Font(bold=True), PatternFill(patternType="solid", fgColor="FFFF00")
    summary = [
        ("直近3期平均 増収率", p(rep.avg_sales_growth_3y), "0.0%"),
        ("直近3期平均 増益率", p(rep.avg_op_growth_3y), "0.0%"),
        ("成長性判定", f"{rep.growth_label}" + (f"（{rep.growth_score_3y:.1f}%）" if rep.growth_score_3y is not None else ""), None),
        (f"EPS 年平均成長率（{rep.eps_cagr_years or '-'}年）", p(rep.eps_cagr), "0.0%"),
        ("EPS と株価の相関係数", rep.eps_price_corr, "0.00"),
        ("過去 PER 中央値", rep.per_median, "0.0"),
        (f"現在 PER（EPS: {rep.current_eps_label or '-'}）", rep.current_per, "0.0"),
        ("PER 中央値から見た目安株価", rep.per_fair_price, "#,##0"),
        ("現在の株価", rep.current_price, "#,##0"),
    ]
    for i, (k, v, fmt) in enumerate(summary):
        ws.cell(row=s + i, column=1, value=k).font = bold
        c = ws.cell(row=s + i, column=4, value=v if v is not None else "-")
        if fmt and v is not None:
            c.number_format = fmt
        if k.startswith("成長性判定"):
            c.fill, c.font = yellow, bold
    ws.column_dimensions["A"].width = 30
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 13

    # 売上高と営業利益の推移
    chart_row = s + len(summary) + 2
    bar = BarChart()
    bar.type, bar.style = "col", 10
    bar.title, bar.y_axis.title, bar.x_axis.title = "売上高と営業利益の推移（百万円）", "百万円", "年度"
    bar.add_data(Reference(ws, min_col=2, max_col=3, min_row=top, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    bar.height, bar.width = 8, 18
    ws.add_chart(bar, f"A{chart_row}")

    # EPS と期末株価の推移（2軸）
    line = LineChart()
    line.title, line.y_axis.title = "EPS と期末株価の推移", "EPS（円）"
    line.add_data(Reference(ws, min_col=9, min_row=top, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    line2 = LineChart()
    line2.add_data(Reference(ws, min_col=10, min_row=top, max_row=last), titles_from_data=True)
    line2.y_axis.axId, line2.y_axis.title = 200, "株価（円）"
    line2.y_axis.crosses = "max"
    line += line2
    line.height, line.width = 8, 18
    ws.add_chart(line, f"G{chart_row}")

    # EPS と株価の散布図
    sc = ScatterChart()
    sc.title, sc.style = "EPS と株価の相関", 13
    sc.x_axis.title, sc.y_axis.title = "EPS", "株価"
    ser = Series(Reference(ws, min_col=10, min_row=top + 1, max_row=last),
                 Reference(ws, min_col=9, min_row=top + 1, max_row=last), title="期末")
    ser.marker = Marker(symbol="circle", size=7)
    ser.graphicalProperties.line.noFill = True
    sc.series.append(ser)
    sc.height, sc.width = 8, 18
    ws.add_chart(sc, f"A{chart_row + 17}")


def build_ir_excel(rep):
    wb = px.Workbook()
    ws = wb.active
    ws.title = "IRレポート"
    write_ir_sheet(ws, rep)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def add_ir_sheet_to_workbook(xlsx_bytes, rep):
    """理論株価の Excel（テンプレート）に IR レポートのシートを追加したバイト列を返す"""
    wb = px.load_workbook(io.BytesIO(xlsx_bytes))
    write_ir_sheet(wb.create_sheet("IRレポート"), rep)
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def ir_file_name(rep):
    name = re.sub(r'[\\/:*?"<>|]', "_", (rep.company or rep.code or "report").strip())
    return f"IR_report_{name}.xlsx"
